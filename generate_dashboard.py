# -*- coding: utf-8 -*-
"""
PICL Dashboard Generator
------------------------
Reads every Excel file inside the PSC and BLDC folders, rebuilds the dashboard data
and writes  Dashboard\\index.html .  Just double-click  Generate_Dashboard.bat .

Folder rule  :  PSC\\Assembly\\PSC-Assembly-Production-Sep.xlsx
File name    :  UNIT-Department-Type-Mon.xlsx   (Type = Production or Reloading)
"""
import sys, os, re, json, shutil, datetime, subprocess, collections

try:
    import openpyxl
except ImportError:
    print("Installing openpyxl (one time only)...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
    import openpyxl

import warnings
warnings.filterwarnings("ignore")

try:
    sys.stdout.reconfigure(errors="replace")
except Exception:
    pass
ROOT = os.path.dirname(os.path.abspath(__file__))
LOG = []
_print = print
def print(*a, **k):  # every message is also saved to Dashboard\generate_log.txt
    _print(*a, **k)
    LOG.append(" ".join(str(x) for x in a))
NOTES = {}
def note(path, msg):
    NOTES.setdefault(path, []).append(msg)
DASH = os.path.join(ROOT, "Dashboard")
TEMPLATE = os.path.join(DASH, "template.html")
OUTPUT = os.path.join(DASH, "index.html")
ARCHIVE = os.path.join(DASH, "Archive")
VERSION = "2026-10-07.1 (auto generator: PSC + BLDC, folder scan)"

DEFAULT_ASM_CATEGORY = {
    "Open(L)": "Rotor Lock", "Open(N)": "Noise Cabin", "Ground/ARC": "Noise Cabin", "HN": "Noise Cabin",
    "BN": "Noise Cabin", "Jam Tight": "Noise Cabin", "PLAY": "Noise Cabin", "TN": "Noise Cabin",
    "Open(T)": "Testing Panel", "Jam Tight2": "Testing Panel", "HV/IR": "Testing Panel", "HW/LW": "Testing Panel",
    "Surge/SHORT": "Testing Panel", "HR/LR": "Testing Panel", "LVS (Tight)": "Testing Panel",
    "Miscellaneous": "Misc"}

# ---------------------------------------------------------------- helpers
def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except Exception:
        return None

def pdate(v):
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    if isinstance(v, str):
        t = v.strip()
        for fmt in ("%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y", "%Y-%m-%d", "%d.%m.%y", "%d-%m-%y", "%d/%m/%y", "%d-%b-%Y", "%d %b %Y", "%d-%b-%y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.datetime.strptime(t, fmt).date().isoformat()
            except ValueError:
                pass
    return None

def clean(v):
    return v.strip() if isinstance(v, str) else ""

def norm_line(v):
    if isinstance(v, (int, float)):
        return int(v)
    t = clean(v)
    return int(t) if t.isdigit() else t

def norm_model(m):
    return clean(m).upper()

def classify(path):
    """Return (unit, dept, kind) from folder / file name, or None."""
    rel = os.path.relpath(path, ROOT).replace("\\", "/")
    low = rel.lower()
    base = os.path.basename(low)
    unit = "BLDC" if re.search(r"\bbldc\b|bldc[-_ ]", low) else ("PSC" if re.search(r"\bpsc\b|psc[-_ ]", low) else None)
    if unit is None:
        return None
    if re.search(r"molding|moulding", base):
        dept = "molding"
    elif "winding" in base:
        dept = "winding"
    elif "shell" in base:
        dept = "shell"
    elif re.search(r"assy|assembly", base):
        dept = "assembly"
    else:
        return None
    kind = "reloading" if re.search(r"reload|reject", base) else "production"
    return unit, dept, kind

# ---------------------------------------------------------------- production
def parse_production(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((ws for ws in wb.worksheets if "daily entry" in ws.title.lower()), None)
    if sheet is None:
        raise RuntimeError("no sheet with 'Daily Entry' in its name. Sheets found: " + ", ".join(w.title for w in wb.worksheets))
    try:
        sheet.reset_dimensions()
    except Exception:
        pass
    out = []
    no_actual = bad_date = bad_row = 0
    all_dates = []
    first_bad = ""
    for n, r in enumerate(sheet.iter_rows(min_row=3, values_only=True), start=3):
        r = list(r) + [None] * 60
        if r[1] in (None, "") or clean(r[1]).lower() in ("date",):
            continue
        d = pdate(r[1])
        if not d:
            bad_date += 1
            first_bad = first_bad or "row %d: '%s'" % (n, r[1])
            continue
        act = num(r[9])
        if act is None:
            no_actual += 1
            continue
        all_dates.append((n, d))
        try:
            slots = []
            for k in range(5):
                b = 20 + 5 * k
                slots.append(dict(t=clean(r[b]), mn=num(r[b + 1]) or 0, mh=num(r[b + 2]) or 0, dp=clean(r[b + 3]), rm=clean(r[b + 4])))
            shifted = slots[0]["mn"] == 0 and slots[1]["mn"] > 0 and bool(slots[0]["t"])
            det = []
            for j in range(5):
                if shifted:
                    if j + 1 >= 5:
                        continue
                    tm, dp = slots[j + 1], slots[j]
                else:
                    tm = dp = slots[j]
                if tm["mn"] > 0 or tm["mh"] > 0:
                    det.append(dict(dept=dp["dp"] or "Others", reason=tm["t"] or "Others", manhr=round(tm["mh"], 4), remark=dp["rm"]))
            g = lambda i, dv=0: (num(r[i]) if num(r[i]) is not None else dv)
            out.append(dict(
                _row=n,
                date=d, line=norm_line(r[2]), shift=clean(r[3]), uphTarget=g(4, None), frame=clean(r[6]), model=clean(r[7]),
                plan=g(8), actual=act, gap=g(12), targetHrs=g(13), actualHrs=g(14), hrsGap=g(15), manpower=g(16),
                targetManpower=g(56, None), injury=g(10), nearMiss=g(11), totalLossHrs=g(18),
                availLossManHr=g(45), perfLossManHr=g(46), defectQty=g(55, None), lossDetails=det))
        except Exception as e:
            bad_row += 1
            first_bad = first_bad or "row %d: %s" % (n, e)
    if no_actual:
        note(path, "%d row(s) have a Date but no Actual qty yet - not shown until Actual is filled" % no_actual)
    if bad_date:
        note(path, "%d row(s) skipped - Date not understood (%s). Use a real date like 05-10-2026" % (bad_date, first_bad))
    if bad_row:
        note(path, "%d row(s) skipped because of an error (%s)" % (bad_row, first_bad))
    if len(all_dates) >= 3:
        D = datetime.date.fromisoformat
        ords = sorted(D(d).toordinal() for _, d in all_dates)
        med = ords[len(ords) // 2]
        fixed, odd = [], []
        by_row = {rec["_row"]: rec for rec in out}
        for n, d in all_dates:
            dd = D(d)
            if abs(dd.toordinal() - med) <= 45:
                continue
            ok = False
            for yr in (datetime.date.fromordinal(med).year,):
                try:
                    cand = dd.replace(year=yr)
                except ValueError:
                    continue
                if abs(cand.toordinal() - med) <= 45 and n in by_row:  # only the YEAR was wrong
                    by_row[n]["date"] = cand.isoformat()
                    fixed.append("row %d: %s -> %s" % (n, dd.strftime("%d.%m.%Y"), cand.strftime("%d.%m.%Y")))
                    ok = True
            if not ok:
                odd.append((n, d))
        if fixed:
            note(path, "%d date(s) had a wrong YEAR and were corrected automatically (%s). Please fix them in the Excel file too." % (len(fixed), "; ".join(fixed[:4])))
        if odd:
            ex = ", ".join("row %d = %s" % (n, D(d).strftime("%d.%m.%Y")) for n, d in odd[:4])
            note(path, "%d row(s) have a date far away from the other dates (%s) - please check the Date column" % (len(odd), ex))
    for rec in out:
        rec.pop("_row", None)
    if not out:
        note(path, "NO DATA ROWS FOUND in sheet '%s' (needs Date in column B and Actual in column J)" % sheet.title)
    return out

# ---------------------------------------------------------------- reloading / rejection
def find_header(rows):
    for i, r in enumerate(rows):
        names = [clean(c).lower().replace("\n", " ") for c in r]
        if "date" in names and ("line" in names or "machine" in names or "model" in names):
            return i
    return None

def parse_reloading(path):
    """Returns (records, category_map). Handles assembly-style and winding-style registers."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    cat_map = {}
    for ws in wb.worksheets:
        if ws.title.lower().startswith("master"):
            rows = list(ws.iter_rows(min_row=1, max_row=200, values_only=True))
            hdr = [clean(c).lower() for c in rows[0]]
            if "defect group" in hdr and "defect name" in hdr:
                gi, ni = hdr.index("defect group"), hdr.index("defect name")
                for r in rows[1:]:
                    if len(r) > max(gi, ni) and r[ni] and r[gi]:
                        cat_map[clean(r[ni])] = "Misc" if clean(r[gi]).lower().startswith("misc") else clean(r[gi])
    recs = []
    for ws in wb.worksheets:
        try:
            ws.reset_dimensions()
        except Exception:
            pass
        rows = [list(r) for r in ws.iter_rows(min_row=1, max_row=12, values_only=True)]
        h = find_header(rows)
        if h is None:
            continue
        head = [clean(c).replace("\n", " ") for c in rows[h]]
        low = [x.lower() for x in head]
        for alias in ("machine", "line no", "line no.", "machine no", "machine no."):
            if alias in low and "line" not in low:
                low[low.index(alias)] = "line"
        ci = {n: low.index(n) for n in ("date", "line", "shift", "frame", "model") if n in low}
        if "line" not in ci or "model" not in ci:
            continue
        tot_i = next((i for i, x in enumerate(low) if x.startswith("total rej")), None)
        reason_cols = {}
        winding_style = any(re.match(r"reason \d+", x) for x in low)
        if winding_style:
            above = [clean(c) for c in rows[h - 1]] if h > 0 else []
            for i, x in enumerate(low):
                if re.match(r"reason \d+", x):
                    lab = above[i] if i < len(above) and above[i] else head[i]
                    m = re.search(r"\((.*)\)", lab)
                    reason_cols[i] = (m.group(1).strip() if m else lab) or head[i]
        else:
            start = ci["model"] + 1
            end = tot_i if tot_i is not None else len(head)
            for i in range(start, end):
                if head[i]:
                    reason_cols[i] = head[i]
        for r in ws.iter_rows(min_row=h + 2, values_only=True):
            r = list(r) + [None] * 70
            d = pdate(r[ci["date"]])
            if not d or not clean(r[ci["model"]]):
                continue
            reasons = {}
            for i, name in reason_cols.items():
                v = num(r[i])
                if v:
                    reasons[name] = int(v) if float(v).is_integer() else v
            total = num(r[tot_i]) if tot_i is not None else None
            if total is None:
                total = sum(reasons.values())
            if not total:
                continue
            cats = {}
            if not winding_style:
                for name, q in reasons.items():
                    c = cat_map.get(name) or DEFAULT_ASM_CATEGORY.get(name) or "Visual Stage"
                    cats[c] = cats.get(c, 0) + q
            recs.append(dict(date=d, line=norm_line(r[ci["line"]]), shift=clean(r[ci.get("shift", 0)]) if "shift" in ci else "",
                             frame=clean(r[ci["frame"]]) if "frame" in ci else "", model=clean(r[ci["model"]]),
                             total=int(total) if float(total).is_integer() else total, categories=cats, reasons=reasons))
        break
    return recs, cat_map

# ---------------------------------------------------------------- build
def main():
    print("=" * 60)
    print(" PICL Dashboard Generator")
    print("=" * 60)
    if not os.path.exists(TEMPLATE):
        print("ERROR: Dashboard\\template.html not found.")
        return 1
    prod = collections.defaultdict(list)
    rel = collections.defaultdict(list)
    cat_map = {}
    files = []
    for dp, dn, fn in os.walk(ROOT):
        if os.path.abspath(dp).startswith(DASH):
            continue
        for f in fn:
            if f.lower().endswith((".xlsx", ".xlsm")) and not f.startswith("~$"):
                c = classify(os.path.join(dp, f))
                if c:
                    files.append((c, os.path.join(dp, f)))
    if not files:
        print("No Excel files found in the PSC / BLDC folders.")
        return 1
    failed = []
    for (unit, dept, kind), path in sorted(files, key=lambda x: x[1]):
        name = os.path.relpath(path, ROOT)
        try:
            if kind == "production":
                rows = parse_production(path)
                prod[(unit, dept)].extend(rows)
            else:
                rows, cm = parse_reloading(path)
                rel[(unit, dept)].extend(rows)
                if unit == "PSC" and dept == "assembly":
                    cat_map.update(cm)
            print(f"  OK  {name:<60} {len(rows):>5} rows")
            for m in NOTES.get(path, []):
                print("      NOTE:", m)
        except Exception as e:
            failed.append(name)
            print(f"  !!  {name}  ->  {e}")
            print("      (If this file is open in Excel, save and CLOSE it, then run again.)")
    # link defects
    data = {}
    for unit in ("PSC", "BLDC"):
        ud = {}
        for dept in ("assembly", "winding", "molding", "shell"):
            P = prod.get((unit, dept), [])
            R = rel.get((unit, dept), [])
            seen = set()
            P2 = []
            for r in P:  # drop exact duplicate rows (same file added twice)
                k = (r["date"], r["line"], r["shift"], r["model"], r["plan"], r["actual"], r["manpower"])
                if k in seen:
                    continue
                seen.add(k)
                P2.append(r)
            P = sorted(P2, key=lambda r: (r["date"], str(r["line"]), r["shift"]))
            if R:
                tot = collections.defaultdict(float)
                days = set()
                for x in R:
                    tot[(x["date"], x["line"], x["shift"], norm_model(x["model"]))] += x["total"]
                    days.add((x["date"], x["line"], x["shift"]))
                for r in P:
                    if r["defectQty"] is None:  # not filled by the sheet's own formula -> join with reloading register
                        k = (r["date"], r["line"], r["shift"], norm_model(r["model"]))
                        if k in tot:
                            v = tot[k]
                            r["defectQty"] = int(v) if float(v).is_integer() else v
                        elif (r["date"], r["line"], r["shift"]) in days:
                            r["defectQty"] = 0
            else:  # no reloading register: build rejection rows from the Defect column of the production sheet
                R = [dict(date=r["date"], line=r["line"], shift=r["shift"], frame=r["frame"], model=r["model"],
                          total=int(r["defectQty"]), categories={}, reasons={}) for r in P if r["defectQty"]]
            ud[dept] = P
            ud[dept + "Rejection"] = sorted(R, key=lambda x: (x["date"], str(x["line"])))
        data[unit] = ud
    warns = []
    for unit in ("PSC", "BLDC"):
        for dept in ("assembly", "winding", "molding"):
            if rel.get((unit, dept)):
                pd_ = {r["date"] for r in data[unit][dept]}
                rd_ = {r["date"] for r in data[unit][dept + "Rejection"]}
                miss = sorted(pd_ - rd_)[-3:]
                last_rel = max(rd_) if rd_ else ""
                late = [d for d in sorted(pd_) if d > last_rel]
                if late:
                    warns.append(f"{unit} {dept.capitalize()}: no Reloading data after {last_rel or 'start'} (production exists till {late[-1]}). Add the latest Reloading file.")
    psc = {k: data["PSC"][k] for k in ("assembly", "winding", "assemblyRejection", "windingRejection")}
    if data["PSC"]["shell"]:
        psc["shell"], psc["shellRejection"] = data["PSC"]["shell"], data["PSC"]["shellRejection"]
    psc["meta"] = {"assemblyReasonCategory": cat_map or DEFAULT_ASM_CATEGORY, "windingReasonCategory": {},
                   "lastUpdated": datetime.date.today().isoformat(), "scriptVersion": VERSION}
    bldc = {k: v for k, v in data["BLDC"].items() if v or not k.endswith("Rejection")}
    html = open(TEMPLATE, encoding="utf-8").read()
    def put(html, tag_id, obj):
        js = json.dumps(obj, separators=(",", ":"), ensure_ascii=False).replace("</", "<\\/")
        pat = re.compile(r'(<script id="%s" type="application/json">)(.*?)(</script>)' % tag_id, re.S)
        if not pat.search(html):
            raise RuntimeError("template is missing the %s block" % tag_id)
        return pat.sub(lambda m: m.group(1) + js + m.group(3), html, count=1)
    html = put(html, "dashboard-data", psc)
    html = put(html, "bldc-data", bldc)
    if os.path.exists(OUTPUT):
        with open(OUTPUT, encoding="utf-8") as fh:
            if fh.read() == html:
                print("-" * 60)
                print("  No new data - dashboard is already up to date.")
                return 0
    with open(OUTPUT, "w", encoding="utf-8") as fh:
        fh.write(html)
    print("-" * 60)
    for unit, ud in (("PSC", psc), ("BLDC", bldc)):
        for dept in ("assembly", "winding", "molding", "shell"):
            rows = ud.get(dept, [])
            if rows or (unit == "PSC" and dept in ("assembly", "winding")) or (unit == "BLDC" and dept != "shell"):
                ds = sorted({r["date"] for r in rows})
                rng = f"{ds[0]}  to  {ds[-1]}  ({len(ds)} days)" if ds else "no data"
                print(f"  {unit:<5}{dept.capitalize():<10}{len(rows):>6} rows   {rng}")
    print("-" * 60)
    for f in failed:
        warns.append("NOT READ: " + f + "  - data from this file is missing in the dashboard")
    for w in warns:
        print("  WARNING:", w)
    print("  Dashboard ready:", OUTPUT)
    return 0

def write_log():
    try:
        os.makedirs(DASH, exist_ok=True)
        with open(os.path.join(DASH, "generate_log.txt"), "w", encoding="utf-8") as fh:
            fh.write("\n".join(LOG))
    except Exception:
        pass

if __name__ == "__main__":
    try:
        code = main()
    except Exception:
        import traceback
        print("\nERROR - the dashboard could not be generated:")
        print(traceback.format_exc())
        code = 1
    write_log()
    if code:
        print("A copy of these messages is saved in Dashboard\\generate_log.txt")
    sys.exit(code)
